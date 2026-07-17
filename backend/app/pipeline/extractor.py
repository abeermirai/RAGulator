from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document as DocxDocument
from openpyxl import load_workbook


@dataclass
class ExtractedPage:
    page_number: int
    text: str


@dataclass
class ExtractedDocument:
    pages: list[ExtractedPage]
    full_text: str


def extract_pdf(path: Path) -> ExtractedDocument:
    pages: list[ExtractedPage] = []
    with fitz.open(path) as doc:
        for idx, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if text:
                pages.append(ExtractedPage(page_number=idx, text=text))
    full_text = "\n\n".join(p.text for p in pages)
    return ExtractedDocument(pages=pages, full_text=full_text)


def extract_docx(path: Path) -> ExtractedDocument:
    doc = DocxDocument(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)
    return ExtractedDocument(pages=[ExtractedPage(page_number=1, text=text)], full_text=text)


def extract_xlsx(path: Path) -> ExtractedDocument:
    wb = load_workbook(path, read_only=True, data_only=True)
    chunks: list[str] = []
    for sheet in wb.worksheets:
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            chunks.append(f"Sheet: {sheet.title}\n" + "\n".join(rows))
    text = "\n\n".join(chunks)
    return ExtractedDocument(pages=[ExtractedPage(page_number=1, text=text)], full_text=text)


def extract_text(path: Path) -> ExtractedDocument:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix in {".docx", ".doc"}:
        return extract_docx(path)
    if suffix in {".xlsx", ".xls"}:
        return extract_xlsx(path)
    if suffix in {".txt", ".md", ".csv"}:
        text = path.read_text(encoding="utf-8", errors="ignore").strip()
        return ExtractedDocument(pages=[ExtractedPage(page_number=1, text=text)], full_text=text)
    raise ValueError(f"Unsupported file type: {suffix}")
